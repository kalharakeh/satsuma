"""v1.1 tools: fit under a size, video/audio extras, PDF pages/watermark/tables/sign,
photo sorting, duplicates, AI upscale, and blurring faces or regions."""
from __future__ import annotations

import hashlib
import io
import math
import os
import re
import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

from engine import (ANY, GIF_VF, H264, Action, ConversionError, Param, Progress, _fitz, _open_pdf, _save_pdf,
                    _same_format, _stem, beside, ffmpeg_exe, ffmpeg_exec, flatten, kind_of, media_duration,
                    natural_key, open_image, register, run_ffmpeg, save_image, scale_short, staged, staged_dir, unique)
import features as F  # attributes are looked up at call time (import cycle with engine)

MB = 1024 * 1024


def models_dir() -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / "models"


def has_audio(src: Path) -> bool:
    import subprocess
    from engine import NOWIN
    r = subprocess.run([ffmpeg_exe(), "-hide_banner", "-i", str(src)], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", creationflags=NOWIN)
    return "Audio:" in r.stderr


def video_size(src: Path) -> tuple[int, int]:
    import subprocess
    from engine import NOWIN
    r = subprocess.run([ffmpeg_exe(), "-hide_banner", "-i", str(src)], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", creationflags=NOWIN)
    m = re.search(r"Video:.*?(\d{2,5})x(\d{2,5})", r.stderr)
    if not m:
        raise ConversionError("No video stream found.")
    w, h = int(m[1]), int(m[2])
    if re.search(r"rotate\s*:\s*-?(90|270)", r.stderr) or "displaymatrix: rotation of -90" in r.stderr \
            or "displaymatrix: rotation of 90" in r.stderr:
        w, h = h, w
    return w, h


def span(progress: Progress, a: float, b: float) -> Progress:
    return lambda f: progress(a + (b - a) * max(0.0, min(f, 1.0)))


# =========================================================== fit under X MB

SIZE_PARAM = Param("mb", "Target size (MB)", "choice", "10", "the result will be at most this big",
                   choices=("8", "10", "16", "20", "25", "50", "100"))


def _target(params: dict) -> int:
    try:
        mb = float(str(params.get("mb") or "10").replace(",", ".").lower().replace("mb", "").strip())
    except ValueError:
        raise ConversionError("Enter the target size as a number of MB, e.g. 10.")
    if mb < 0.01:
        raise ConversionError("That target is too small - 10 KB is the minimum.")
    return int(mb * MB)


def _mb_tag(target: int) -> str:
    mb = target / MB
    return f"{mb:g}MB" if mb >= 1 else f"{round(mb * 1024)}KB"


def fit_image(src: Path, target: int, progress: Progress) -> Path:
    from PIL import Image
    im = open_image(src)
    if getattr(im, "is_animated", False):
        im.seek(0)
        im = im.copy()
    alpha = "A" in im.getbands() or (im.mode == "P" and "transparency" in im.info)
    fmt, ext = ("WEBP", ".webp") if alpha or src.suffix.lower() == ".webp" else ("JPEG", ".jpg")
    icc = im.info.get("icc_profile")
    im.info = {"icc_profile": icc} if icc else {}

    def encode(img, q) -> bytes:
        buf = io.BytesIO()
        save_image(img, buf, fmt, q)
        return buf.getvalue()

    work, best = im, None
    for step in range(12):
        lo, hi, found = 35, 95, None
        while lo <= hi:  # highest quality that fits
            q = (lo + hi) // 2
            data = encode(work, q)
            if len(data) <= target:
                found, lo = data, q + 1
            else:
                hi = q - 1
        progress(min(0.1 + step * 0.08, 0.9))
        if found:
            best = found
            break
        # still too big at quality 35: shrink the picture and try again
        size = len(encode(work, 35))
        s = max(0.3, min(0.9, math.sqrt(target / size) * 0.95))
        work = work.resize((max(1, int(work.width * s)), max(1, int(work.height * s))), Image.Resampling.LANCZOS)
    if best is None:
        raise ConversionError("Couldn't get this image that small.")
    dest = beside(src, ext, _mb_tag(target))
    with staged(dest) as tmp:
        tmp.write_bytes(best)
    return dest


def fit_av(src: Path, target: int, progress: Progress) -> Path:
    dur = media_duration(ffmpeg_exe(), src)
    if not dur:
        raise ConversionError("Can't read the length of this file.")
    if kind_of(src) == "audio":
        kbps = target * 8 / 1000 / dur * 0.96
        if kbps < 16:
            raise ConversionError(f"Too long to fit in {_mb_tag(target)} - even at the lowest quality.")
        return run_ffmpeg(src, beside(src, ".mp3", _mb_tag(target)),
                          ["-vn", "-map", "0:a:0", "-c:a", "libmp3lame", "-b:a", f"{int(min(320, kbps))}k"], progress)
    audio = has_audio(src)
    w, h = video_size(src)
    dest = beside(src, ".mp4", _mb_tag(target))
    budget = target
    for attempt in range(3):  # x264 can overshoot a little: retry with a smaller budget
        total_kbps = budget * 8 / 1000 / dur * 0.96
        a_kbps = 0 if not audio else (128 if total_kbps > 2000 else 96 if total_kbps > 600 else 64 if total_kbps > 250 else 32)
        v_kbps = int(total_kbps - a_kbps)
        if v_kbps < 60:
            raise ConversionError(f"Too long to fit in {_mb_tag(target)} as video - try a bigger target or trim it first.")
        short = min(w, h)
        cap = next(lim for lim, need in ((1080, 3500), (720, 1500), (540, 800), (480, 450), (360, 0)) if v_kbps >= need)
        vf = ["-vf", scale_short(cap)] if cap < short else []
        a = ["-c:a", "aac", "-b:a", f"{a_kbps}k"] if audio else ["-an"]
        p = span(progress, attempt / 3 if attempt else 0, 1)
        with tempfile.TemporaryDirectory(prefix="satsuma-2pass-") as td:
            log = str(Path(td) / "pass")
            common = ["-c:v", "libx264", "-preset", "medium", "-b:v", f"{v_kbps}k", "-pix_fmt", "yuv420p",
                      *vf, "-passlogfile", log]
            ffmpeg_exec(["-i", str(src), *common, "-pass", "1", "-an", "-f", "null", os.devnull], dur, span(p, 0, 0.45))
            with staged(dest) as tmp:
                ffmpeg_exec(["-i", str(src), *common, "-pass", "2", *a, "-movflags", "+faststart", str(tmp)], dur,
                            span(p, 0.45, 1))
        if dest.stat().st_size <= target:
            return dest
        budget = int(budget * target / dest.stat().st_size * 0.95)
        dest.unlink()
    raise ConversionError(f"Couldn't get this video under {_mb_tag(target)}.")


def fit_pdf(src: Path, target: int, progress: Progress) -> Path:
    if src.stat().st_size <= target:
        raise ConversionError(f"Already under {_mb_tag(target)}.")
    fz = _fitz()
    data = None
    steps = [(150, 75), (120, 65), (100, 55), (85, 45), (72, 40), (60, 35), (50, 30)]
    for i, (dpi, q) in enumerate(steps):
        doc = _open_pdf(src)
        if hasattr(doc, "rewrite_images"):
            doc.rewrite_images(dpi_threshold=dpi + 10, dpi_target=dpi, quality=q, set_to_gray=False)
        buf = io.BytesIO()
        try:
            doc.save(buf, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True, use_objstms=1)
        except TypeError:
            doc.save(buf, garbage=4, deflate=True, clean=True)
        data = buf.getvalue()
        progress((i + 1) / len(steps))
        if len(data) <= target:
            break
    if data is None or len(data) > target:
        raise ConversionError(f"Got it down to {len(data) / MB:.1f} MB, but not under {_mb_tag(target)} - "
                              "this PDF is mostly text and fonts.")
    dest = beside(src, ".pdf", _mb_tag(target))
    with staged(dest) as tmp:
        tmp.write_bytes(data)
    return dest


def fit_size(src: Path, progress: Progress, params: dict):
    target = _target(params)
    k = kind_of(src)
    if k == "image":
        return fit_image(src, target, progress)
    if k in ("video", "audio"):
        return fit_av(src, target, progress)
    if k == "pdf":
        return fit_pdf(src, target, progress)
    raise ConversionError("Fit-to-size works for images, video, audio and PDFs.")


register(Action("fit-size", "Fit under a size (MB)…", "Fit MB", "Tools",
                frozenset({"image", "video", "audio", "pdf"}), fit_size, hint="email/upload",
                params=(SIZE_PARAM,), prio=9))


# ================================================================ video

def speed(src: Path, progress: Progress, params: dict):
    try:
        f = float(str(params.get("factor") or "2").lower().replace("x", "").replace(",", "."))
    except ValueError:
        raise ConversionError("Speed must be a number like 2 or 0.5.")
    if not 0.1 <= f <= 16:
        raise ConversionError("Choose a speed between 0.1x and 16x.")
    tempo, chain = f, []
    while tempo > 2.0:
        chain.append("atempo=2.0")
        tempo /= 2.0
    while tempo < 0.5:
        chain.append("atempo=0.5")
        tempo /= 0.5
    chain.append(f"atempo={tempo:.4f}")
    video = kind_of(src) == "video"
    dur = media_duration(ffmpeg_exe(), src) / f
    tag = f"{f:g}x"
    if video:
        a = ["-af", ",".join(chain), "-c:a", "aac", "-b:a", "160k"] if has_audio(src) else ["-an"]
        args = ["-vf", f"setpts=PTS/{f}", *H264[:6], "-crf", "20", *a, "-movflags", "+faststart"]
        return run_ffmpeg(src, beside(src, ".mp4", tag), args, progress, duration=dur)
    ext = src.suffix.lower() if src.suffix.lower() in (".mp3", ".wav", ".flac", ".m4a", ".ogg", ".opus") else ".mp3"
    from engine import AUDIO_CODEC
    return run_ffmpeg(src, beside(src, ext, tag), ["-vn", "-af", ",".join(chain), *AUDIO_CODEC.get(ext, [])],
                      progress, duration=dur)


SPEED_PARAM = Param("factor", "Speed", "choice", "2", "e.g. 2 = twice as fast, 0.5 = slow motion",
                    choices=("0.25", "0.5", "0.75", "1.25", "1.5", "2", "4", "8"))
register(Action("vid-speed", "Change speed…", "Speed", "Tools", frozenset({"video"}), speed, hint="faster/slower",
                params=(SPEED_PARAM,), prio=14))
register(Action("aud-speed", "Change speed…", "Speed", "Tools", frozenset({"audio"}), speed, hint="faster/slower",
                params=(SPEED_PARAM,), prio=15))


def merge_videos(srcs: list[Path], progress: Progress):
    srcs = sorted(srcs, key=natural_key)
    exe = ffmpeg_exe()
    w, h = video_size(srcs[0])
    w, h = w - w % 2, h - h % 2
    inputs, filters, labels, total = [], [], [], 0.0
    n = 0
    for i, s in enumerate(srcs):
        d = media_duration(exe, s)
        total += d
        inputs += ["-i", str(s)]
        filters.append(f"[{i}:v]scale={w}:{h}:force_original_aspect_ratio=decrease,pad={w}:{h}:(ow-iw)/2:(oh-ih)/2,"
                       f"setsar=1,fps=30,format=yuv420p[v{i}]")
        if has_audio(s):
            filters.append(f"[{i}:a]aresample=48000,aformat=channel_layouts=stereo[a{i}]")
        else:
            filters.append(f"anullsrc=r=48000:cl=stereo,atrim=duration={d:.3f}[a{i}]")
        labels.append(f"[v{i}][a{i}]")
        n += 1
    filters.append(f"{''.join(labels)}concat=n={n}:v=1:a=1[v][a]")
    dest = beside(srcs[0], ".mp4", "merged")
    with staged(dest) as tmp:
        ffmpeg_exec([*inputs, "-filter_complex", ";".join(filters), "-map", "[v]", "-map", "[a]",
                     "-c:v", "libx264", "-preset", "medium", "-crf", "20", "-c:a", "aac", "-b:a", "160k",
                     "-movflags", "+faststart", str(tmp)], total, progress)
    return dest


register(Action("vid-merge", "Join into one video", "Join", "Tools", frozenset({"video"}), merge_videos,
                each=False, min_files=2, hint="in name order", prio=8))


def frames(src: Path, progress: Progress, params: dict):
    try:
        every = F.parse_time(params.get("every") or "1")
    except ValueError:
        raise ConversionError("Enter seconds between frames, e.g. 1 or 0.5.")
    if every <= 0:
        raise ConversionError("Seconds between frames must be more than 0.")
    dur = media_duration(ffmpeg_exe(), src)
    if dur and dur / every > 5000:
        raise ConversionError(f"That's {int(dur / every):,} pictures - pick a longer gap.")
    dest = beside(src, "", "frames")
    with staged_dir(dest) as tmp:
        ffmpeg_exec(["-i", str(src), "-vf", f"fps=1/{every}", "-q:v", "2", str(tmp / f"{_stem(src)}-%05d.jpg")],
                    dur, progress)
        if not any(tmp.iterdir()):
            raise ConversionError("No frames came out.")
    return dest


register(Action("vid-frames", "Save frames as pictures…", "Frames", "Tools", frozenset({"video"}), frames,
                hint="JPG every N s", params=(Param("every", "Seconds between frames", "choice", "1",
                                                     choices=("0.5", "1", "2", "5", "10", "30")),), prio=33))


def gif_clip(src: Path, progress: Progress, params: dict):
    exe = ffmpeg_exe()
    total = media_duration(exe, src)
    try:
        start = F.parse_time(params.get("start") or "0")
        end = F.parse_time(params["end"]) if str(params.get("end") or "").strip() else min(total, start + 5)
    except ValueError:
        raise ConversionError("Times must look like 1:30 or 00:01:30.")
    end = min(end, total) if total else end
    if end <= start:
        raise ConversionError("The end time must be after the start time.")
    if end - start > 60:
        raise ConversionError("GIFs longer than a minute get huge - use MP4 for that.")
    width = str(params.get("width") or "480")
    vf = GIF_VF.replace("min(640,iw)", f"min({int(width) if width.isdigit() else 480},iw)")
    return run_ffmpeg(src, beside(src, ".gif", "clip"), ["-t", f"{end - start:.3f}", "-vf", vf, "-loop", "0"],
                      progress, pre=["-ss", f"{start:.3f}"], duration=end - start)


register(Action("vid-gifclip", "GIF from part of the video…", "GIF clip", "Tools", frozenset({"video"}), gif_clip,
                hint="pick a range", prio=16, params=(
                    Param("start", "Start", "time", "0:00", "e.g. 0:15"),
                    Param("end", "End", "time", "", "empty = 5 seconds"),
                    Param("width", "Width (px)", "choice", "480", choices=("320", "480", "640", "800")))))


def burn_subs(src: Path, progress: Progress):
    srt = src.with_suffix(".srt")
    made = []
    if not srt.exists():
        outs = F.transcribe(src, span(progress, 0, 0.5))
        srt = outs[0]
        made = outs
        p2 = span(progress, 0.5, 1)
    else:
        p2 = progress
    dest = beside(src, ".mp4", "subtitled")
    with tempfile.TemporaryDirectory(prefix="satsuma-subs-") as td:
        shutil.copy(srt, Path(td) / "subs.srt")  # the subtitles filter can't take Windows paths with ':' easily
        style = "FontName=Arial,FontSize=18,PrimaryColour=&H00FFFFFF,OutlineColour=&H80000000,BorderStyle=1,Outline=2,Shadow=0,MarginV=22"
        with staged(dest) as tmp:
            a = ["-c:a", "copy"] if has_audio(src) else ["-an"]
            ffmpeg_exec(["-i", str(src), "-vf", f"subtitles=subs.srt:force_style='{style}'", *H264[:6], "-crf", "20",
                         *a, "-movflags", "+faststart", str(tmp)], media_duration(ffmpeg_exe(), src), p2, cwd=td)
    return [dest, *made]


register(Action("vid-burn", "Burn in subtitles", "Captions", "Tools", frozenset({"video"}), burn_subs,
                hint="subtitles on video", prio=13))


DENOISE = "highpass=f=70,lowpass=f=14000,afftdn=nr=14:nf=-30:tn=1"


def denoise(src: Path, progress: Progress):
    if kind_of(src) == "video":
        if not has_audio(src):
            raise ConversionError("This video has no sound.")
        return run_ffmpeg(src, beside(src, src.suffix.lower() if src.suffix.lower() in (".mp4", ".mov", ".mkv") else ".mp4",
                                      "clean-audio"),
                          ["-map", "0:v?", "-map", "0:a:0", "-c:v", "copy", "-af", DENOISE, "-c:a", "aac", "-b:a", "192k"],
                          progress)
    from engine import AUDIO_CODEC
    ext = src.suffix.lower() if src.suffix.lower() in AUDIO_CODEC else ".mp3"
    return run_ffmpeg(src, beside(src, ext, "denoised"), ["-vn", "-map", "0:a:0", "-af", DENOISE, *AUDIO_CODEC[ext]],
                      progress)


register(Action("aud-denoise", "Remove background noise", "Denoise", "Tools", frozenset({"audio"}), denoise,
                hint="hiss & hum", prio=12))
register(Action("vid-denoise", "Remove background noise", "Denoise", "Tools", frozenset({"video"}), denoise,
                hint="clean audio", prio=24))


# ================================================================== PDF

def pdf_pages(src: Path, progress: Progress, params: dict):
    """params: order = [page indexes to keep, in order], rotate = {index: degrees}."""
    doc = _open_pdf(src)
    order = [int(i) for i in params.get("order", [])]
    if not order:
        raise ConversionError("You removed every page - nothing to save.")
    if any(i < 0 or i >= doc.page_count for i in order):
        raise ConversionError("Page list doesn't match this PDF.")
    rot = {int(k): int(v) for k, v in (params.get("rotate") or {}).items()}
    for i, deg in rot.items():
        if deg % 360:
            p = doc[i]
            p.set_rotation((p.rotation + deg) % 360)
    doc.select(order)
    dest = beside(src, ".pdf", "edited")
    with staged(dest) as tmp:
        _save_pdf(doc, tmp)
    return dest


register(Action("pdf-pages", "Reorder / delete / rotate pages…", "Pages", "Tools", frozenset({"pdf"}), pdf_pages,
                hint="edit pages", editor="pdfpages", prio=9))


def watermark(src: Path, progress: Progress, params: dict):
    text = (params.get("text") or "").strip()
    if not text:
        raise ConversionError("Enter the watermark text.")
    try:
        opacity = max(0.05, min(1.0, float(str(params.get("opacity") or "15").strip("% ")) / 100))
    except ValueError:
        opacity = 0.15
    fz = _fitz()
    doc = _open_pdf(src)
    for n, page in enumerate(doc):
        r = page.rect
        diag = math.hypot(r.width, r.height)
        size = max(12, min(110, diag * 0.62 / max(len(text), 1) / 0.55))
        width = fz.get_text_length(text, fontname="helv", fontsize=size)
        center = fz.Point(r.width / 2, r.height / 2)
        angle = math.degrees(math.atan2(r.height, r.width))
        start = fz.Point(center.x - width / 2, center.y + size / 3)
        page.insert_text(start, text, fontsize=size, fontname="helv", color=(0.45, 0.45, 0.45),
                         fill_opacity=opacity, stroke_opacity=opacity, overlay=True,
                         morph=(center, fz.Matrix(angle)), rotate=0)
        progress((n + 1) / doc.page_count)
    dest = beside(src, ".pdf", "watermarked")
    with staged(dest) as tmp:
        _save_pdf(doc, tmp)
    return dest


register(Action("pdf-watermark", "Add a watermark…", "Watermark", "Tools", frozenset({"pdf"}), watermark,
                hint="text across pages", prio=32, params=(
                    Param("text", "Text", "choice", "CONFIDENTIAL", choices=("CONFIDENTIAL", "DRAFT", "COPY", "SAMPLE")),
                    Param("opacity", "Opacity (%)", "choice", "15", choices=("10", "15", "25", "40")))))


def pdf_tables(src: Path, progress: Progress):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    doc = _open_pdf(src)
    wb = Workbook()
    wb.remove(wb.active)
    count = 0
    for n, page in enumerate(doc):
        try:
            tabs = page.find_tables()
        except Exception:
            tabs = []
        for k, t in enumerate(tabs.tables if hasattr(tabs, "tables") else tabs):
            rows = t.extract()
            if not rows or len(rows) < 2:
                continue
            count += 1
            ws = wb.create_sheet(f"p{n + 1}" + (f"-{k + 1}" if k else ""))
            for row in rows:
                ws.append([_num(c) for c in row])
            for col in range(1, max(len(r) for r in rows) + 1):
                width = max(len(str(r[col - 1] or "")) for r in rows if len(r) >= col)
                ws.column_dimensions[get_column_letter(col)].width = min(60, max(8, width + 2))
        progress((n + 1) / doc.page_count)
    if not count:
        raise ConversionError("No tables found. (Scanned PDFs need OCR first.)")
    dest = beside(src, ".xlsx", "tables")
    with staged(dest) as tmp:
        wb.save(tmp)
    return dest


_NUMRE = re.compile(r"^-?\d{1,3}([ ,. ]\d{3})*([.,]\d+)?$|^-?\d+([.,]\d+)?$")


def _num(v):
    if v is None:
        return None
    s = str(v).replace("\n", " ").strip()
    if _NUMRE.match(s):
        t = s.replace(" ", "").replace(" ", "")
        if "," in t and "." in t:
            t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
        elif "," in t:
            t = t.replace(",", ".") if len(t.split(",")[-1]) != 3 else t.replace(",", "")
        try:
            f = float(t)
            return int(f) if f.is_integer() and "." not in t else f
        except ValueError:
            pass
    return s


register(Action("pdf-tables", "Tables to Excel", "Tables", "Convert", frozenset({"pdf"}), pdf_tables,
                hint="XLSX", prio=60))


def pdf_sign(src: Path, progress: Progress, params: dict):
    """params: image (signature PNG/JPG), page (index), rect [x0, y0, x1, y1] in PDF points."""
    img = Path(params.get("image") or "")
    if not img.is_file():
        raise ConversionError("Signature picture not found.")
    fz = _fitz()
    doc = _open_pdf(src)
    i = int(params.get("page", 0))
    if not 0 <= i < doc.page_count:
        raise ConversionError("That page doesn't exist.")
    x0, y0, x1, y1 = (float(v) for v in params["rect"])
    # transparent background for signatures photographed or scanned on white paper
    from PIL import Image
    with Image.open(img) as s:
        s = s.convert("RGBA")
        if "A" not in Image.open(img).getbands():
            gray = s.convert("L")
            alpha = gray.point(lambda v: 0 if v > 200 else min(255, (230 - v) * 3))
            s.putalpha(alpha)
        buf = io.BytesIO()
        s.save(buf, "PNG")
    doc[i].insert_image(fz.Rect(x0, y0, x1, y1), stream=buf.getvalue(), keep_proportion=True, overlay=True)
    dest = beside(src, ".pdf", "signed")
    with staged(dest) as tmp:
        _save_pdf(doc, tmp)
    return dest


register(Action("pdf-sign", "Sign…", "Sign", "Tools", frozenset({"pdf"}), pdf_sign, hint="place signature",
                editor="sign", prio=13))


# ============================================================ photos & files

MEDIA = {"image", "video"}


def _media_files(srcs: list[Path]) -> list[Path]:
    out = []
    for s in srcs:
        if s.is_dir():
            out += [p for p in sorted(s.iterdir()) if p.is_file() and kind_of(p) in MEDIA]
        elif kind_of(s) in MEDIA:
            out.append(s)
    return out


def sort_by_date(srcs: list[Path], progress: Progress):
    files = _media_files(srcs)
    if not files:
        raise ConversionError("No photos or videos to sort.")
    moves = []
    for n, f in enumerate(files):
        d = F._taken(f)
        folder = f.parent / d.strftime("%Y") / d.strftime("%Y-%m")
        folder.mkdir(parents=True, exist_ok=True)
        dest = unique(folder / f.name)
        os.replace(f, dest)
        moves.append((f, dest))
        progress((n + 1) / len(files))
    folders = sorted({str(b.parent) for _, b in moves})
    return [Path(p) for p in folders], {"rename": [[str(b), str(a)] for a, b in moves]}


register(Action("sort-date", "Sort into folders by date", "By date", "Tools", frozenset({"image", "video", "folder"}),
                sort_by_date, each=False, hint="YYYY / YYYY-MM", prio=70))


def _hash(p: Path) -> str:
    h = hashlib.blake2b(digest_size=20)
    with open(p, "rb") as f:
        while chunk := f.read(1 << 20):
            h.update(chunk)
    return h.hexdigest()


def duplicates(srcs: list[Path], progress: Progress):
    files: list[Path] = []
    for s in srcs:
        files += [p for p in s.rglob("*") if p.is_file()] if s.is_dir() else [s]
    files = [f for f in files if "Duplicates" not in f.parts and not f.name.startswith("~")]
    by_size = defaultdict(list)
    for f in files:
        try:
            by_size[f.stat().st_size].append(f)
        except OSError:
            pass
    groups, candidates = [], [g for size, g in by_size.items() if len(g) > 1 and size > 0]
    total = sum(len(g) for g in candidates) or 1
    done = 0
    for g in candidates:
        by_hash = defaultdict(list)
        for f in g:
            by_hash[_hash(f)].append(f)
            done += 1
            progress(0.9 * done / total)
        groups += [v for v in by_hash.values() if len(v) > 1]
    if not groups:
        raise ConversionError(f"No duplicates among {len(files):,} files.")
    root = srcs[0] if srcs[0].is_dir() else srcs[0].parent
    dup_dir = unique(root / "Duplicates")
    dup_dir.mkdir()
    moves, report = [], [f"Satsuma duplicate report - {len(groups)} groups\n"]
    for g in groups:
        g.sort(key=lambda p: (p.stat().st_mtime, len(p.name)))  # keep the oldest, shortest name
        keep, extra = g[0], g[1:]
        report.append(f"KEPT   {keep}")
        for e in extra:
            dest = unique(dup_dir / e.name)
            os.replace(e, dest)
            moves.append((e, dest))
            report.append(f"moved  {e}  ->  {dest.name}")
        report.append("")
    (dup_dir / "_report.txt").write_text("\n".join(report), encoding="utf-8")
    return [dup_dir], {"rename": [[str(b), str(a)] for a, b in moves]}


register(Action("find-dupes", "Find duplicates (moves extras aside)", "Dupes", "Tools", frozenset({"folder"}),
                duplicates, each=False, hint="same content", prio=71))


# ================================================================== AI upscale

_esr = None


def upscale(src: Path, progress: Progress, params: dict):
    import numpy as np
    from PIL import Image
    import aipack
    aipack.ensure(progress, (0.0, 0.2))
    global _esr
    if _esr is None:
        import onnxruntime as ort
        opts = ort.SessionOptions()
        opts.log_severity_level = 3
        _esr = ort.InferenceSession(str(models_dir() / "realesr-general-x4v3.onnx"), opts,
                                    providers=["CPUExecutionProvider"])
    factor = 2 if str(params.get("scale") or "2").strip().lower().rstrip("x") == "2" else 4
    im = open_image(src)
    if getattr(im, "is_animated", False):
        raise ConversionError("Animated images aren't supported.")
    if max(im.size) * factor > 12000:
        raise ConversionError(f"The result would be over 12000 px - this picture is already big enough.")
    alpha = im.getchannel("A") if "A" in im.getbands() else None
    rgb = np.asarray(im.convert("RGB"), dtype=np.float32) / 255.0
    h, w, _ = rgb.shape
    tile, pad = 192, 12
    out = np.zeros((h * 4, w * 4, 3), dtype=np.float32)
    ys, xs = list(range(0, h, tile)), list(range(0, w, tile))
    n, total = 0, len(ys) * len(xs)
    for y in ys:
        for x in xs:
            y0, x0 = max(y - pad, 0), max(x - pad, 0)
            y1, x1 = min(y + tile + pad, h), min(x + tile + pad, w)
            inp = rgb[y0:y1, x0:x1].transpose(2, 0, 1)[None]
            res = _esr.run(None, {"input": np.ascontiguousarray(inp)})[0][0].transpose(1, 2, 0)
            ty, tx = (y - y0) * 4, (x - x0) * 4
            th, tw = min(tile, h - y) * 4, min(tile, w - x) * 4
            out[y * 4:y * 4 + th, x * 4:x * 4 + tw] = res[ty:ty + th, tx:tx + tw]
            n += 1
            progress(0.2 + 0.75 * n / total)
    big = Image.fromarray((np.clip(out, 0, 1) * 255 + 0.5).astype(np.uint8), "RGB")
    if factor == 2:
        big = big.resize((w * 2, h * 2), Image.Resampling.LANCZOS)
    if alpha is not None:
        big.putalpha(alpha.resize(big.size, Image.Resampling.LANCZOS))
    fmt, ext = _same_format(src)
    if fmt in ("GIF", "ICO", "BMP"):
        fmt, ext = "PNG", ".png"
    dest = beside(src, ext, f"{factor}x")
    big.info = {"icc_profile": im.info["icc_profile"]} if im.info.get("icc_profile") else {}
    with staged(dest) as tmp:
        save_image(big, tmp, fmt, 95 if fmt in ("JPEG", "WEBP") else None)
    return dest


register(Action("img-upscale", "Upscale with AI…", "Upscale", "Tools", frozenset({"image"}), upscale,
                hint="2x / 4x sharper", prio=13,
                params=(Param("scale", "Enlarge", "choice", "2", choices=("2", "4")),)))


# ================================================================ blur faces / regions

def detect_faces(path: Path) -> list[list[int]]:
    """Face boxes [x, y, w, h] in image pixels (YuNet, offline). Empty if none or unavailable."""
    try:
        import cv2
        import numpy as np
    except ImportError:
        return []
    im = flatten(open_image(path))
    w, h = im.size
    s = min(1.0, 1600 / max(w, h))
    small = im.resize((max(1, int(w * s)), max(1, int(h * s)))) if s < 1 else im
    arr = cv2.cvtColor(np.asarray(small.convert("RGB")), cv2.COLOR_RGB2BGR)
    det = cv2.FaceDetectorYN.create(str(models_dir() / "yunet.onnx"), "", (arr.shape[1], arr.shape[0]), 0.6, 0.3, 5000)
    det.setInputSize((arr.shape[1], arr.shape[0]))
    _, faces = det.detect(arr)
    boxes = []
    for f in (faces if faces is not None else []):
        x, y, bw, bh = (float(v) / s for v in f[:4])
        m = 0.18  # a bit of margin so hair/chin are covered too
        boxes.append([int(x - bw * m), int(y - bh * m * 1.4), int(bw * (1 + 2 * m)), int(bh * (1 + 2.4 * m))])
    return boxes


def blur(src: Path, progress: Progress, params: dict):
    from PIL import Image, ImageDraw, ImageFilter
    rects = params.get("rects") or []
    if not rects:
        raise ConversionError("Nothing marked to blur.")
    style = params.get("style", "blur")
    im = open_image(src)
    base = im.convert("RGBA") if "A" in im.getbands() else im.convert("RGB")
    W, H = base.size
    for x, y, w, h in rects:
        x0, y0 = max(0, int(x)), max(0, int(y))
        x1, y1 = min(W, int(x + w)), min(H, int(y + h))
        if x1 - x0 < 2 or y1 - y0 < 2:
            continue
        region = base.crop((x0, y0, x1, y1))
        if style == "pixelate":
            k = max(6, min(x1 - x0, y1 - y0) // 12)
            region = region.resize((max(1, (x1 - x0) // k), max(1, (y1 - y0) // k)), Image.Resampling.BILINEAR)
            region = region.resize((x1 - x0, y1 - y0), Image.Resampling.NEAREST)
        elif style == "black":
            region = Image.new(base.mode, region.size, (0, 0, 0, 255) if base.mode == "RGBA" else (0, 0, 0))
        else:
            r = max(8, min(x1 - x0, y1 - y0) // 5)
            region = region.filter(ImageFilter.GaussianBlur(r)).filter(ImageFilter.GaussianBlur(r // 2))
        if style in ("blur", "pixelate"):  # soft oval for faces, sharp box for black bars
            mask = Image.new("L", region.size, 0)
            ImageDraw.Draw(mask).rounded_rectangle([0, 0, region.width - 1, region.height - 1],
                                                   radius=min(region.size) // 4, fill=255)
            base.paste(region, (x0, y0), mask)
        else:
            base.paste(region, (x0, y0))
    fmt, ext = _same_format(src)
    if fmt in ("GIF", "ICO"):
        fmt, ext = "PNG", ".png"
    base.info = {}  # blurring is for privacy: drop EXIF/GPS too
    dest = beside(src, ext, "blurred")
    with staged(dest) as tmp:
        save_image(base, tmp, fmt, 92 if fmt in ("JPEG", "WEBP") else None)
    return dest


register(Action("img-blur", "Blur faces & areas…", "Blur", "Tools", frozenset({"image"}), blur,
                hint="privacy", editor="blur", prio=14))


# ================================================================ app-only tiles

def _app_only(*_a, **_k):
    raise ConversionError("This tool runs inside the Satsuma app.")


register(Action("send-phone", "Send to phone (QR code)", "Phone", "Tools", ANY, _app_only, each=False,
                hint="over Wi-Fi", ui="phone", prio=6))


def _install_ai(files, progress: Progress):
    import aipack
    aipack.ensure(progress, (0.0, 1.0))
    return aipack.pack_dir()


register(Action("ai-install", "Download AI pack", "AI pack", "Hidden", frozenset(), _install_ai, each=False,
                min_files=0))
